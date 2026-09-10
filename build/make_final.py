# -*- coding: utf-8 -*-
"""
Final SIH 2026 deck for SIH26056 (VIMAAN), Team Bharat Bytes.

Difference from make_visual.py: every AI-generated illustration is gone.
  * data visuals  -> real matplotlib plates from plates.py (n_*.png)
  * structural diagrams (pipeline, risk->response, the CPI chain)
    -> native PowerPoint shapes, drawn here, so they stay crisp and editable
  * no stock people, no robots, no decorative aircraft, no shield icons

Template chrome (footer bar, SIH logo, slide numbers, team oval, titles) is
left exactly as shipped. Six slides, title slide included.

    python plates.py && python make_final.py
"""
import os
from PIL import Image
from _prims import (Presentation, Inches, Pt, RGBColor, PP_ALIGN, MSO_ANCHOR,
                    MSO_SHAPE, TPL, TEAM, HERE,
                    NAVY, INK, MUTED, FAINT, AMBER, TEAL, RED, WHITE,
                    BLUE_FILL, BLUE_LINE, AMBER_FILL, AMBER_LINE,
                    GREY_FILL, GREY_LINE,
                    noshadow, tb, par, run, box, eyebrow, card, chip, pic, find, drop)

OUT = os.path.join(HERE, "SIH2026-VIMAAN-SIH26056-BharatBytes.pptx")

GREEN = RGBColor(0x1B, 0x7A, 0x3C)
GOLD  = RGBColor(0xF3, 0xC6, 0x7A)

X0, X1 = 0.35, 12.98
BW  = X1 - X0
COL = 6.10
XR  = X0 + COL + 0.18          # 6.63


# ---- plate placement --------------------------------------------------------
def plate(slide, name, x, y, w, frame=True):
    """Place a generated chart, deriving height from the file so nothing squashes."""
    p = os.path.join(HERE, name)
    iw, ih = Image.open(p).size
    h = w * ih / float(iw)
    if frame:
        rule(slide, x, y + h + 0.035, w, thick=0.5)
    sh = slide.shapes.add_picture(p, Inches(x), Inches(y), width=Inches(w))
    sh.height = Inches(h)
    noshadow(sh)
    return y + h


WALKTHROUGH_URL = ("https://drive.google.com/drive/folders/"
                   "1umIfrbjnLhSdqTqecGn315W5UIARY4um")
LIVE_URL = "https://vimaan-console.vercel.app"
REPO_URL = "https://github.com/hersheysss3/vimaan-airfare-index"


def fix_hyperlink_theme(prs):
    """Make hyperlinks obey the deck palette.

    PowerPoint paints any hyperlinked run with the theme's <a:hlink> colour,
    which overrides the colour set on the run. The default is a bright blue
    that appears nowhere else in this deck. Repoint the theme instead.
    """
    from pptx.oxml.ns import qn
    hexes = ("1F497D", "1F497D")           # navy for new and followed links
    for master in prs.slide_masters:
        theme = master.part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/"
            "relationships/theme")
        root = theme._element if hasattr(theme, "_element") else None
        if root is None:
            from lxml import etree
            root = etree.fromstring(theme.blob)
        for tag, hx in zip(("hlink", "folHlink"), hexes):
            for el in root.iter(qn("a:" + tag)):
                for child in list(el):
                    el.remove(child)
                srgb = el.makeelement(qn("a:srgbClr"), {"val": hx})
                el.append(srgb)
        from lxml import etree
        theme._blob = etree.tostring(root, xml_declaration=True,
                                     encoding="UTF-8", standalone=True)


def link_run(paragraph, text, url, size=8.5, color=None, bold=False):
    """A run that is actually clickable in PowerPoint and in the exported PDF.

    Printing a URL as plain text looks identical on screen and does nothing
    when clicked, which is worse than useless in a deck read as a PDF.
    """
    r = run(paragraph, text, size, color if color is not None else NAVY,
            bold=bold)
    r.hyperlink.address = url
    return r


def caption(slide, x, y, w, text, size=7.5):
    tf = tb(slide, x, y, w, 0.24)
    p = par(tf, first=True, line=1.06)
    run(p, text, size, FAINT, italic=True)


def exhibit(slide, x, y, w, num, pointer):
    """Rule, exhibit number, then the mandated SIH pointer on the same line.
    One label carries both, so nothing is said twice."""
    rule(slide, x, y, w, thick=1.6, color=NAVY)
    tf = tb(slide, x, y + 0.055, w, 0.22)
    pp = par(tf, first=True, line=1.02)
    run(pp, "EXHIBIT %d" % num, 8, AMBER, bold=True)
    run(pp, "     " + pointer.upper(), 8, NAVY, bold=True)
    return y + 0.28


def source(slide, x, y, w, text):
    """The line that says where a number came from. Statistical releases carry
    one under every table; generated decks never do."""
    tf = tb(slide, x, y, w, 0.22)
    pp = par(tf, first=True, line=1.04)
    run(pp, "Source:  ", 6.8, NAVY, bold=True)
    run(pp, text, 6.8, FAINT)
    return y + 0.20


def note_block(slide, x, y, w, label, text, h=None, accent=None):
    """Hairline, small-caps label, text. Replaces every filled card."""
    rule(slide, x, y, w, thick=1.2, color=accent or NAVY)
    tf = tb(slide, x, y + 0.05, w, 0.20)
    pp = par(tf, first=True, line=1.02)
    run(pp, label.upper(), 7.6, accent or NAVY, bold=True)
    tf = tb(slide, x, y + 0.245, w, (h or 0.70) - 0.245)
    pp = par(tf, first=True, line=1.10)
    run(pp, text, 7.8, MUTED)
    return y + (h or 0.70)


def kpi(slide, x, y, w, value, label, vsize=13):
    tf = tb(slide, x, y, w, 0.30)
    p = par(tf, first=True)
    run(p, value, vsize, AMBER, bold=True)
    tf = tb(slide, x, y + 0.24, w, 0.20)
    p = par(tf, first=True)
    run(p, label, 6.5, FAINT, bold=True)


# ---- native diagram primitives ---------------------------------------------
def rule(slide, x, y, w, thick=1.0, color=None):
    """A hairline. Rules instead of card fills is what stops this reading as AI."""
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(thick / 72.0))
    sh.fill.solid(); sh.fill.fore_color.rgb = color or GREY_LINE
    sh.line.fill.background(); noshadow(sh)
    return sh


def vrule(slide, x, y, h, color=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                Inches(1.0 / 72.0), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color or GREY_LINE
    sh.line.fill.background(); noshadow(sh)
    return sh


def stage_row(slide, x, y, w, h, stages):
    """The five processing stages as a ruled table, not a row of filled boxes."""
    n = len(stages)
    cw = w / n
    rule(slide, x, y, w, thick=1.6, color=NAVY)
    for i, (num, name, sub) in enumerate(stages):
        cx = x + i * cw
        if i:
            vrule(slide, cx, y + 0.02, h - 0.04)
        tf = tb(slide, cx + 0.09, y + 0.07, cw - 0.18, 0.20)
        pp = par(tf, first=True, line=1.02)
        run(pp, num + "   ", 8, AMBER, bold=True)
        run(pp, name, 8.8, NAVY, bold=True)
        tf = tb(slide, cx + 0.09, y + 0.29, cw - 0.18, h - 0.34)
        pp = par(tf, first=True, line=1.08)
        run(pp, sub, 7.6, MUTED)
    rule(slide, x, y + h, w, thick=0.75)



def spec_table(slide, x, y, w, rows, lw=1.66, rh=0.246):
    """Parameter / value table with hairlines. Reads as a datasheet, not a card."""
    rule(slide, x, y, w, thick=1.6, color=NAVY)
    yy = y + 0.05
    for k, v in rows:
        tf = tb(slide, x, yy, lw, rh - 0.02)
        pp = par(tf, first=True, line=1.04)
        run(pp, k, 7.4, NAVY, bold=True)
        tf = tb(slide, x + lw, yy, w - lw, rh - 0.02)
        pp = par(tf, first=True, line=1.04)
        run(pp, v, 7.8, INK)
        yy += rh
        rule(slide, x, yy - 0.035, w, thick=0.5)
    return yy


def map_rows(slide, x, y, w, rows, rh=0.30, lw=2.58):
    """Two-column 'risk -> designed response' mapping. Replaces the shield icons."""
    for i, (left, right) in enumerate(rows):
        yy = y + i * rh
        tf = tb(slide, x, yy, lw, rh - 0.02, anchor=MSO_ANCHOR.MIDDLE)
        p = par(tf, first=True, line=1.04)
        run(p, left, 7.8, INK, bold=True)
        a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                   Inches(x + lw + 0.04), Inches(yy + rh / 2 - 0.05),
                                   Inches(0.16), Inches(0.10))
        a.fill.solid(); a.fill.fore_color.rgb = AMBER
        a.line.fill.background(); noshadow(a)
        tf = tb(slide, x + lw + 0.26, yy, w - lw - 0.26, rh - 0.02, anchor=MSO_ANCHOR.MIDDLE)
        p = par(tf, first=True, line=1.04)
        run(p, right, 7.8, MUTED)


def chain(slide, x, y, w, h, steps):
    """Fares -> index -> CPI -> policy. Ruled stages, matching the process table
    on slide 3; the final stage is marked by an amber rule rather than a fill."""
    n = len(steps)
    cw = w / n
    AW, AH = 0.15, 0.105          # arrow size
    NAME_Y = y + 0.07             # top of the stage-name line
    NAME_H = 0.22
    rule(slide, x, y, w, thick=1.6, color=NAVY)
    for i, (head, sub) in enumerate(steps):
        cx = x + i * cw
        last = (i == n - 1)
        if last:
            rule(slide, cx, y, cw, thick=1.6, color=AMBER)
        # a flow separates with arrows, not with rules — using both made the
        # arrows straddle the dividers, so the vertical rules are gone here
        tf = tb(slide, cx + 0.09, NAME_Y, cw - 0.30, NAME_H)
        pp = par(tf, first=True, line=1.02)
        run(pp, head, 8.6, AMBER if last else NAVY, bold=True)
        tf = tb(slide, cx + 0.09, y + 0.28, cw - 0.30, h - 0.32)
        pp = par(tf, first=True, line=1.06)
        run(pp, sub, 7.2, MUTED)
        if i < n - 1:
            # centred on the stage boundary, and on the name line's optical
            # centre so the eye reads name -> arrow -> name in one sweep
            a = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(cx + cw - AW / 2.0),
                Inches(NAME_Y + NAME_H / 2.0 - AH / 2.0),
                Inches(AW), Inches(AH))
            a.fill.solid(); a.fill.fore_color.rgb = AMBER
            a.line.fill.background(); noshadow(a)
    rule(slide, x, y + h, w, thick=0.5)


def set_title(slide, text, left=None, width=None, size=None, align=PP_ALIGN.CENTER):
    t = slide.shapes.title
    if t is None:
        for sh in slide.shapes:
            if sh.name.startswith("Title"):
                t = sh; break
    tf = t.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    for r in list(p.runs)[1:]:
        r._r.getparent().remove(r._r)
    if p.runs:
        p.runs[0].text = text
        if size:
            p.runs[0].font.size = Pt(size)
    else:
        r = p.add_run(); r.text = text
        r.font.name = "Times New Roman"; r.font.bold = True
        r.font.size = Pt(size or 36)
    p.alignment = align
    if left is not None:
        t.left = Inches(left)
    if width is not None:
        t.width = Inches(width)
    t.top = Inches(0.04); t.height = Inches(1.10)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return t


# =============================================================================
prs = Presentation(TPL)
sld = prs.slides._sldIdLst
for sid in list(sld)[6:]:
    rId = sid.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rId); sld.remove(sid)
S = list(prs.slides)

for s in S[1:]:
    for sh in s.shapes:
        if sh.name.startswith("Oval"):
            p = sh.text_frame.paragraphs[0]
            for r in list(p.runs)[1:]:
                r._r.getparent().remove(r._r)
            if p.runs:
                p.runs[0].text = TEAM
                p.runs[0].font.size = Pt(11); p.runs[0].font.bold = True
            p.alignment = PP_ALIGN.CENTER


# =============================================================== SLIDE 1
s = S[0]
drop(find(s, "Subtitle 3"))
TXW = 6.10
tx = find(s, "TextBox 9"); tf = tx.text_frame; tf.clear()
tx.top, tx.left, tx.width, tx.height = Inches(1.22), Inches(0.36), Inches(TXW), Inches(2.30)

ROWS = [("Problem Statement ID", "SIH26056"),
        ("Problem Statement Title",
         "Development of a Real-time Airfare Price Index for India through Automated Web "
         "Scraping of Airline and Online Travel Aggregator Portals for Augmentation of the "
         "Consumer Price Index (CPI)"),
        ("Theme", "Travel & Tourism"),
        ("PS Category", "Software"),
        ("Team Name", TEAM)]
for i, (k, v) in enumerate(ROWS):
    p = par(tf, first=(i == 0), after=6, line=1.06)
    run(p, "•  " + k + " – ", 14, INK, bold=True)
    run(p, v, 13.5, NAVY if k != "Team Name" else AMBER,
        bold=(k in ("Problem Statement ID", "Team Name")))

b = box(s, 0.36, 3.62, TXW, 1.08, fill=NAVY, line=NAVY)
tf = b.text_frame
p = tf.paragraphs[0]
run(p, "VIMAAN", 22, WHITE, bold=True)
run(p, "   Validated Index of Market Airfares", 10.5, RGBColor(0xC8, 0xD6, 0xE8))
p2 = par(tf, line=1.06)
run(p2, "for Augmenting National-statistics", 10.5, RGBColor(0xC8, 0xD6, 0xE8))
p3 = par(tf, line=1.08)
run(p3, "Skyscanner tells you what a ticket costs. VIMAAN tells the RBI whether to raise rates.",
    9.5, GOLD, italic=True)

# the index itself, not an airport photograph
yh = plate(s, "n_hero_index.png", 0.36, 4.78, 4.34, frame=False)
source(s, 0.36, yh + 0.04, 4.34,
       "Modelled on the VIMAAN sampling design. Base 2024 = 100, as in CPI 2024.")
for i, (v, l) in enumerate([("517", "ROUTES, 3× A DAY"),
                            ("78 → 517", "VS DGCA TODAY"),
                            ("SDMX", "FEED INTO CPI 2024")]):
    kpi(s, 4.84, 4.92 + i * 0.52, 1.66, v, l)

# The deck is read as a PDF more often than it is presented. Whoever is
# holding it should be one click away from the prototype running.
rule(s, 0.36, 6.66, TXW, thick=1.6, color=NAVY)
tf = tb(s, 0.36, 6.73, TXW, 0.26)
p = par(tf, first=True, line=1.06)
run(p, "WALKTHROUGH  ", 7.6, AMBER, bold=True)
link_run(p, "Prototype video and screens on Google Drive",
         WALKTHROUGH_URL, size=9, bold=True)
tf = tb(s, 0.36, 7.02, TXW, 0.26)
p = par(tf, first=True, line=1.06)
run(p, "LIVE  ", 7.2, AMBER, bold=True)
link_run(p, "vimaan-console.vercel.app", LIVE_URL, size=8)
run(p, "        SOURCE  ", 7.2, AMBER, bold=True)
link_run(p, "github.com/hersheysss3/vimaan-airfare-index", REPO_URL, size=8)


# =============================================================== SLIDE 2
s = S[1]
drop(find(s, "TextBox 8"))
set_title(s, "PROPOSED SOLUTION", left=1.85, width=8.70, size=30)

b = box(s, X0, 1.20, BW, 0.42, fill=NAVY, line=NAVY)
tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
run(p, "VIMAAN   ", 11, GOLD, bold=True)
run(p, "Validated Index of Market Airfares for Augmenting National-statistics.   ",
    9.5, WHITE, bold=True)
run(p, "An automated replacement for the manual airfare price collection that "
       "currently supplies the CPI transport group.", 9.5, RGBColor(0xC8, 0xD6, 0xE8))

y2 = exhibit(s, X0, 1.68, COL, 2, "Detailed explanation of the proposed solution")
y3 = exhibit(s, XR, 1.68, COL, 3, "How it addresses the problem")
yl = plate(s, "n_today.png",   X0, y2, COL, frame=False)
yr = plate(s, "n_surface.png", XR, y3, COL, frame=False)
source(s, X0, yl + 0.02, COL,
       "Route count from DGCA Monthly Domestic Traffic Statistics. "
       "Present coverage per PIB PRID 1810467.")
source(s, XR, yr + 0.02, COL,
       "Fare grid modelled on the VIMAAN sampling design, Delhi–Mumbai, "
       "September 2026 departures.")

eyebrow(s, X0, 5.86, "Innovation and uniqueness of the solution", w=6.30)
INNO = [("01", "Full price surface sampled",
         "Eight advance-purchase windows across every departure date in a 60-day horizon, "
         "each fare attributed to its month of departure."),
        ("02", "Multilateral index estimators",
         "Jevons elementary aggregates, then GEKS-Jevons on a 25-month rolling window, "
         "which removes chain drift."),
        ("03", "Statutory collection basis",
         "Rule 135(2) tariff pages and licensed APIs supply two of three lanes. "
         "robots.txt observed; no personal data."),
        ("04", "Reproducible output",
         "Hashed raw snapshots, a methodology endpoint for each value, and SDMX-JSON "
         "for direct ingestion by MoSPI.")]
cw = (BW - 3 * 0.20) / 4.0
rule(s, X0, 6.08, BW, thick=1.6, color=NAVY)
for i, (n, h, t) in enumerate(INNO):
    cx = X0 + i * (cw + 0.20)
    if i:
        vrule(s, cx - 0.10, 6.11, 0.74)
    tf = tb(s, cx, 6.135, cw, 0.22)
    p = par(tf, first=True, line=1.02)
    run(p, n + "   ", 8, AMBER, bold=True)
    run(p, h, 8.6, NAVY, bold=True)
    tf = tb(s, cx, 6.355, cw, 0.50)
    p = par(tf, first=True, line=1.10)
    run(p, t, 7.6, MUTED)


# =============================================================== SLIDE 3
s = S[2]
drop(find(s, "TextBox 8"))
set_title(s, "TECHNICAL APPROACH", left=1.85, width=8.70, size=30)

LW = 7.42
RX = X0 + LW + 0.16
RW = X1 - RX

eyebrow(s, X0, 1.24, "Methodology and process for implementation", w=LW)
stage_row(s, X0, 1.46, LW, 0.74, [
    ("1", "COLLECT",  "statutory tariff pages, licensed APIs, public portals"),
    ("2", "CLEAN",    "IATA codes, one currency, all-in fares, outliers out"),
    ("3", "ADJUST",   "hedonic model strips stops, duration, cabin, baggage"),
    ("4", "INDEX",    "Jevons → GEKS–Jevons → DGCA passenger weights"),
    ("5", "PUBLISH",  "REST + SDMX-JSON, dashboard, surge alerts to DGCA"),
])
caption(s, X0, 2.26, LW,
        "The pipeline runs without manual intervention. Sampling parameters are given in the specification, right.",
        size=7.6)

y4 = exhibit(s, X0, 2.56, LW, 4, "Effect of the index formula on the published figure")
ydr = plate(s, "n_drift.png", X0, y4, 6.40, frame=False)
ysr = source(s, X0, ydr + 0.02, LW,
             "Simulation with constant true prices. Estimators per ONS Methodology "
             "Working Paper 12 and the ILO/IMF CPI Manual 2020.")
tf = tb(s, X0, ysr + 0.02, LW, 0.24)
p = par(tf, first=True, line=1.04)
run(p, "P_J(0,t) = Πⱼ ( pⱼᵗ / pⱼ⁰ )^(1/n)          "
       "Index(t) = Σᵢ wᵢ Iᵢ(t) / Σᵢ wᵢ × 100,  wᵢ = DGCA base-period passengers",
    7.6, INK)

eyebrow(s, RX, 1.24, "Technologies to be used", w=RW)
STACK = [("Collect", "Python 3.11 · Playwright · httpx · Redis + RQ"),
         ("Store", "PostgreSQL 16 + TimescaleDB · Parquet / DuckDB"),
         ("Process", "Polars · NumPy / SciPy · Great Expectations"),
         ("Statistics", "statsmodels · X-13ARIMA-SEATS · index library"),
         ("Serve", "FastAPI · Pydantic v2 · SDMX-JSON · Swagger"),
         ("Interface", "React 18 + Vite · visx · MapLibre + deck.gl"),
         ("Operate", "Docker Compose · Prometheus + Grafana · CI")]
yy = 1.46
for k, v in STACK:
    tf = tb(s, RX, yy, RW, 0.26)
    p = par(tf, first=True, line=1.05)
    run(p, k.upper() + "   ", 7.5, AMBER, bold=True)
    run(p, v, 8, INK)
    yy += 0.265

eyebrow(s, RX, 3.36, "Index specification", w=RW)
yspec = spec_table(s, RX, 3.58, RW, [
    ("Collection",   "3 times daily, 02:00 / 11:00 / 19:00 IST"),
    ("Lead buckets", "L1, L3, L7, L14, L21, L30, L45, L60"),
    ("Horizon",      "60 days forward, every departure date"),
    ("Attribution",  "month of departure (acquisition basis in parallel)"),
    ("Elementary",   "Jevons geometric mean, per route and bucket"),
    ("Multilateral", "GEKS-Jevons, cross-checked with Time-Product-Dummy"),
    ("Window",       "25 months, rolling, extended by mean splicing"),
    ("Weights",      "DGCA base-period route passengers"),
    ("Seasonal",     "X-13ARIMA-SEATS, Indian moving-holiday regressors"),
    ("Output",       "REST, SDMX-JSON, monthly bulletin"),
])

rule(s, RX, yspec + 0.10, RW, thick=1.6, color=NAVY)
tf = tb(s, RX, yspec + 0.18, RW, 0.24)
p = par(tf, first=True, line=1.04)
run(p, "OUT OF SCOPE BY DESIGN   ", 7.4, AMBER, bold=True)
tf = tb(s, RX, yspec + 0.38, RW, 0.50)
p = par(tf, first=True, line=1.10)
run(p, "No CAPTCHA solving, no login or paywall circumvention, no fingerprint "
       "spoofing, and no personal data at any stage. Two of the three intake "
       "lanes do not scrape.", 7.8, MUTED)


# =============================================================== SLIDE 4
s = S[3]
drop(find(s, "TextBox 8"))
set_title(s, "FEASIBILITY AND VIABILITY", left=1.85, width=8.70, size=30)

eyebrow(s, X0, 1.24, "Analysis of the feasibility of the idea", w=6.30)
FEAS = [("Collection basis is statutory",
         "Rule 135(2), Aircraft Rules 1937 obliges every airline to publish its tariff. Two of "
         "our three intake lanes therefore need no scraping at all."),
        ("Components are standard and unlicensed",
         "Python, Playwright, PostgreSQL, statsmodels, FastAPI, React, Docker. Nothing exotic, "
         "nothing licensed, nothing to procure.")]
yy = 1.44
for h, t in FEAS:
    yy = note_block(s, X0, yy, COL, h, t, h=0.74) + 0.10

y5 = exhibit(s, X0, 3.16, COL, 5, "Implied data volume")
yv = plate(s, "n_volume.png", X0, y5, COL, frame=False)
source(s, X0, yv + 0.02, COL,
       "Derived from the sampling design set out in the specification on slide 3.")

y6 = exhibit(s, XR, 1.24, COL, 6, "Potential challenges and risks")
yrk = plate(s, "n_risk.png", XR, y6, COL, frame=False)
source(s, XR, yrk + 0.02, COL,
       "Team assessment. Each entry is answered in the strategies below.")

rule(s, XR, 4.92, COL, thick=1.6, color=NAVY)
tf = tb(s, XR, 4.975, COL, 0.22)
p = par(tf, first=True, line=1.02)
run(p, "STRATEGIES FOR OVERCOMING THESE CHALLENGES", 8, NAVY, bold=True)
map_rows(s, XR, 5.22, COL, [
    ("1  Portal blocks us",     "three independent lanes; A and B are statutory and contractual"),
    ("2  Legal challenge",      "Rule 135(2) mandate + Collection of Statistics Act 2008; zero PII"),
    ("3  No 2024 base fares",   "backcast from DGCA yields; published provisional with a band"),
    ("4  Chain drift",          "GEKS–Jevons and TPD on a 25-month rolling window, mean spliced"),
    ("5  Listed ≠ paid",   "yield calibration from DGCA route revenue, published with uncertainty"),
    ("6  Parser breaks",        "config-driven selectors, schema-drift detector, per-lane SLO alarms"),
    ("7  Demo network dies",    "seeded 21-day dataset, fully offline Docker Compose, backup video"),
], rh=0.232, lw=1.52)


# =============================================================== SLIDE 5
s = S[4]
drop(find(s, "TextBox 8"))
set_title(s, "IMPACT AND BENEFITS", left=1.85, width=8.70, size=30)

y7 = exhibit(s, X0, 1.22, COL, 7, "Potential impact on the target audience")
yi = plate(s, "n_impact.png", X0, y7, 5.60, frame=False)
source(s, X0, yi + 0.02, COL,
       "Present method per PIB PRID 1810467. VIMAAN figures are design parameters.")

y8 = exhibit(s, X0, 4.36, COL, 8, "Position within the CPI 2024 basket")
yc = plate(s, "n_cpi.png", X0, y8, 5.40, frame=False)
source(s, X0, yc + 0.02, COL,
       "MoSPI, Consumer Price Index base 2024 = 100. First release, PIB PRID 2227012.")

eyebrow(s, XR, 1.24, "Transmission of the index into policy", w=6.30)
chain(s, XR, 1.46, COL, 0.76, [
    ("Millions of fares", "517 routes, 3× daily"),
    ("VIMAAN index",      "quality-adjusted"),
    ("CPI 2024",          "transport, 9.43%"),
    ("Repo rate",         "EMIs for 140 crore"),
])

source(s, XR, 2.26, COL,
       "CPI group weight from MoSPI. Repo transmission is the standard monetary "
       "policy channel, not a claim specific to this project.")
rule(s, XR, 2.52, COL, thick=1.6, color=NAVY)
tf = tb(s, XR, 2.575, COL, 0.22)
p = par(tf, first=True, line=1.02)
run(p, "BENEFITS OF THE SOLUTION (SOCIAL, ECONOMIC, ENVIRONMENTAL)", 8, NAVY, bold=True)
BEN = [("Economic — MoSPI and the RBI",
        "An unbiased, quality-adjusted airfare component feeds headline CPI, which sets the "
        "repo rate. Better input, better monetary policy."),
       ("Administrative — DGCA",
        "Tariff monitoring goes from 78 hand-checked routes a month to 517 routes three times "
        "a day, with automatic fare-surge alerts."),
       ("Social — citizens and travellers",
        "Public route and lead-time indices expose the true advance-purchase penalty and "
        "festival-season surges. No such statistic exists for India today."),
       ("Environmental — MoCA",
        "Field price collection is replaced by an automated feed, removing enumerator travel. "
        "The ATF pass-through model shows how much of a fare rise is fuel.")]
yy = 2.84
for h, t in BEN:
    yy = note_block(s, XR, yy, COL, h, t, h=0.82) + 0.04

note_block(s, XR, 6.32, COL, "Statistical — MoSPI, RBI and the IMF",
           "SDMX-JSON is ingested without manual handling. Hashed provenance makes "
           "every published value defensible under statistical audit and under RTI.",
           h=0.54, accent=AMBER)


# =============================================================== SLIDE 6
s = S[5]
drop(find(s, "TextBox 8"))
set_title(s, "RESEARCH AND REFERENCES", left=1.85, width=8.70, size=30)

eyebrow(s, X0, 1.22, "Details / links of the reference and research work  ·  "
                     "each source, and what it settles", w=BW)

CW3 = (BW - 2 * 0.18) / 3.0
CX = [X0, X0 + CW3 + 0.18, X0 + 2 * (CW3 + 0.18)]

BR = [("A", "INDIAN STATUTORY BASIS", NAVY,
       "Establishes that the data is public and that MoSPI may collect it"),
      ("B", "INTERNATIONAL METHOD", AMBER,
       "Establishes that our estimators are the ones NSOs already use"),
      ("C", "OUR OWN VALIDATION", GREEN,
       "Establishes how anyone can check the number we publish")]
for i, (tag, name, col, why) in enumerate(BR):
    bb = box(s, CX[i], 1.44, CW3, 0.36, fill=col, line=col)
    tf = bb.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, tag + "   " + name, 8.5, WHITE, bold=True)
    tf = tb(s, CX[i], 1.84, CW3, 0.30)
    p = par(tf, first=True, line=1.06)
    run(p, why, 7.6, MUTED, italic=True)

A = [("MoSPI — FAQs on the CPI 2024 Series, Annexure V",
      "airfare prices are collected through online platforms", "mospi.gov.in"),
     ("PIB — first CPI release on base 2024 = 100, 12 Feb 2026",
      "358 items · COICOP 2018 · Transport weight 9.43%", "PRID 2227012"),
     ("Aircraft Rules, 1937 — Rule 135(2)",
      "every airline shall publish its established tariff", "civilaviation.gov.in"),
     ("Collection of Statistics Act, 2008",
      "statutory mandate for official price collection", "indiacode.nic.in"),
     ("PIB / MoCA — DGCA Tariff Monitoring Unit",
      "78 routes, monthly, read by hand from airline sites", "PRID 1810467"),
     ("DGCA — Monthly Domestic Traffic Statistics",
      "route-wise passengers — our index weights", "dgca.gov.in"),
     ("MoSPI — Expert Group Report on updation of the CPI",
      "weight derivation and item-basket methodology", "mospi.gov.in")]

B = [("ONS — aggregate index of air fares (FOI-2023-1164)",
      "fixed advance-purchase windows; month-of-departure basis", "ons.gov.uk"),
     ("ONS — Methodology Working Paper 12",
      "index-number methods on web-scraped price data", "ons.gov.uk"),
     ("ONS — multilateral methods in consumer price statistics",
      "GEKS–Jevons, TPD, rolling windows, splicing", "ons.gov.uk"),
     ("ILO / IMF / OECD / Eurostat — CPI Manual 2020",
      "why Jevons, not Carli or Dutot, for elementary aggregates", "imf.org"),
     ("Statistics Canada — Air Transportation Index, 2020",
      "dynamic airline pricing inside an official index", "statcan.gc.ca"),
     ("Statistics Netherlands / Eurostat — web scraping",
      "25-month window guidance and chain-drift evidence", "unece.org"),
     ("IOCL — Aviation Turbine Fuel price circulars",
      "input to the fuel pass-through model", "iocl.com")]

for col, items, num0, accent in ((0, A, 1, NAVY), (1, B, 8, AMBER)):
    yy = 2.22
    for i, (h, t, src) in enumerate(items):
        tf = tb(s, CX[col], yy, CW3, 0.52)
        p = par(tf, first=True, line=1.04)
        run(p, "%d.  " % (num0 + i), 8, accent, bold=True)
        run(p, h, 8, NAVY, bold=True)
        p2 = par(tf, after=0, line=1.04)
        run(p2, "     " + t + "   ", 7.5, MUTED)
        run(p2, src, 7, FAINT, bold=True)
        yy += 0.585

# branch C — our own evidence, and the closing differentiator
note_block(s, CX[2], 2.22, CW3, "Datasets assembled before hour zero",
           "routes_top500.csv (DGCA), dgca_pax_weights.csv, base_2024_fares.csv "
           "(backcast), airports.geojson, atf_prices.csv (IOCL), festival_calendar.csv",
           h=0.86)

note_block(s, CX[2], 3.22, CW3, "How the index will be validated",
           "Benchmarked against MoSPI's own manually collected airfare series over "
           "overlapping months. Cross-lane divergence tests with a flag threshold. "
           "Leave-one-route-out sensitivity on the DGCA weights. A published "
           "revision policy.", h=1.06)

note_block(s, CX[0], 6.04, CW3 * 2 + 0.18, "Standing of the method",
           "The collection basis is statutory and the estimators are those already "
           "in use by the UK Office for National Statistics, Statistics Canada and "
           "Statistics Netherlands on web-collected prices. The contribution here is "
           "their application to Indian route data, with the supporting evidence "
           "published alongside every figure.", h=0.82)

rule(s, CX[2], 4.46, CW3, thick=1.6, color=AMBER)
tf = tb(s, CX[2], 4.515, CW3, 0.22)
p = par(tf, first=True, line=1.02)
run(p, "DISTINCTION FROM A FARE-SEARCH PRODUCT", 8, AMBER, bold=True)
DIST = [("Unit of observation", "a price surface, not a single price"),
        ("Elementary aggregate", "Jevons geometric mean, not an arithmetic average"),
        ("Multilateral stage", "GEKS-Jevons on a rolling window"),
        ("Attribution", "month of departure, DGCA passenger weights"),
        ("Delivery", "an SDMX statistical feed, not only a dashboard"),
        ("Auditability", "a reproducibility hash on every published value"),
        ("Collection basis", "two of three lanes statutory, not scraped")]
yy = 4.76
for k, v in DIST:
    tf = tb(s, CX[2], yy, 1.32, 0.22)
    pp = par(tf, first=True, line=1.02)
    run(pp, k, 7.2, NAVY, bold=True)
    tf = tb(s, CX[2] + 1.36, yy, CW3 - 1.36, 0.22)
    pp = par(tf, first=True, line=1.02)
    run(pp, v, 7.4, MUTED)
    yy += 0.245
    rule(s, CX[2], yy - 0.035, CW3, thick=0.4)

# the evidence column should say where the working thing can be seen
tf = tb(s, CX[2], yy + 0.05, CW3, 0.24)
p = par(tf, first=True, line=1.02)
run(p, "SEE IT RUN  ", 7.2, GREEN, bold=True)
link_run(p, "walkthrough", WALKTHROUGH_URL, size=7.4, bold=True)
run(p, "  ·  ", 7.4, FAINT)
link_run(p, "vimaan-console.vercel.app", LIVE_URL, size=7.4)



# ---- no inherited theme shadows anywhere ------------------------------------
for sl in prs.slides:
    for sh in sl.shapes:
        noshadow(sh)

fix_hyperlink_theme(prs)
prs.save(OUT)
print("saved", OUT, os.path.getsize(OUT) // 1024, "KB")
