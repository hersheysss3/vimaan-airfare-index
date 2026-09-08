# -*- coding: utf-8 -*-
"""
Fills the official SIH 2026 Idea Submission template for SIH26056 (VIMAAN).

Rules honoured:
  * the supplied template file is the base - chrome, fonts, colours, footer bar,
    SIH logo, slide numbers and team oval are left exactly as shipped
  * six slides, title slide included (the template's instruction page is dropped)
  * every guidance pointer from the template is kept as a section eyebrow so a
    reviewer can see each one is answered
  * no paragraphs - points, cards, diagrams, charts
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
# The official SIH template is vendored under build/template/ so the deck can be
# rebuilt from a clean clone. Point SIH_TEMPLATE elsewhere to override.
TPL = os.environ.get("SIH_TEMPLATE") or os.path.join(
    HERE, "template", "SIH2026-IDEA-Presentation-Format.pptx")
OUT = os.path.join(HERE, "SIH2026-VIMAAN-SIH26056-BharatBytes.pptx")

TEAM = "Bharat Bytes"
TEAM_ID = "<TEAM ID>"

# ---- template-native palette -------------------------------------------------
NAVY   = RGBColor(0x1F, 0x49, 0x7D)   # theme dk2
BLUE   = RGBColor(0x00, 0x70, 0xC0)   # footer bar
INK    = RGBColor(0x1A, 0x1F, 0x2B)
MUTED  = RGBColor(0x55, 0x5F, 0x6E)
FAINT  = RGBColor(0x8A, 0x93, 0xA2)
AMBER  = RGBColor(0xB8, 0x6B, 0x00)
TEAL   = RGBColor(0x00, 0x89, 0x7B)
INDIGO = RGBColor(0x5A, 0x4B, 0xC7)
RED    = RGBColor(0xB0, 0x2A, 0x20)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

BLUE_FILL   = RGBColor(0xEA, 0xF1, 0xF9); BLUE_LINE   = RGBColor(0xC2, 0xD6, 0xEB)
AMBER_FILL  = RGBColor(0xFD, 0xF4, 0xE6); AMBER_LINE  = RGBColor(0xE3, 0xBE, 0x7B)
RED_FILL    = RGBColor(0xFD, 0xEC, 0xEA); GREY_FILL   = RGBColor(0xF3, 0xF5, 0xF8)
GREY_LINE   = RGBColor(0xD8, 0xDE, 0xE7)

F = "Arial"

# content band, inside the template chrome
X0, X1 = 0.35, 12.98
Y0, Y1 = 1.28, 6.85


# ---- primitives --------------------------------------------------------------
def noshadow(sh):
    try:
        sh.shadow.inherit = False
    except Exception:
        pass
    return sh


def tb(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    noshadow(box)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.paragraphs[0].alignment = align
    return tf


def par(tf, first=False, after=2, line=None, align=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(after)
    p.space_before = Pt(0)
    if line:
        p.line_spacing = line
    if align is not None:
        p.alignment = align
    elif not first:
        p.alignment = PP_ALIGN.LEFT
    return p


def run(p, text, size=10, color=INK, bold=False, italic=False, font=F):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    return r


def box(slide, x, y, w, h, fill=BLUE_FILL, line=BLUE_LINE, rounded=True, lw=0.75):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shp, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        sh.adjustments[0] = 0.06
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    noshadow(sh)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.10); tf.margin_right = Inches(0.10)
    tf.margin_top = Inches(0.06);  tf.margin_bottom = Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    return sh


def eyebrow(slide, x, y, text, w=6.0, color=NAVY, size=9.5):
    tf = tb(slide, x, y, w, 0.20)
    p = par(tf, first=True)
    run(p, text.upper(), size, color, bold=True)
    return tf


def card(slide, x, y, w, h, head, body, fill=BLUE_FILL, line=BLUE_LINE,
         head_color=NAVY, hsize=9.5, bsize=8.5, kicker=None):
    sh = box(slide, x, y, w, h, fill, line)
    tf = sh.text_frame
    p = tf.paragraphs[0]
    if kicker:
        run(p, kicker + "  ", hsize, AMBER, bold=True)
    run(p, head, hsize, head_color, bold=True)
    if body:
        p2 = par(tf, line=1.08)
        run(p2, body, bsize, MUTED)
    return sh


def icard(slide, x, y, w, h, icon, head, body, fill=BLUE_FILL, line=BLUE_LINE,
          hsize=9.5, bsize=8.5, kicker=None, iso=0.24):
    """Card with a generated line-art icon in a left gutter."""
    sh = box(slide, x, y, w, h, fill, line)
    tf = sh.text_frame
    tf.margin_left = Inches(0.10 + iso + 0.10)
    p = tf.paragraphs[0]
    if kicker:
        run(p, kicker + "  ", hsize, AMBER, bold=True)
    run(p, head, hsize, NAVY, bold=True)
    if body:
        p2 = par(tf, line=1.08)
        run(p2, body, bsize, MUTED)
    pic(slide, "ic_" + icon + ".png", x + 0.10, y + 0.10, h=iso)
    return sh


def chip(slide, x, y, w, h, text, fill=BLUE, color=WHITE, size=7.5, bold=True,
         line=None, rounded=True):
    sh = box(slide, x, y, w, h, fill, line if line else fill, rounded)
    tf = sh.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.04)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, text, size, color, bold=bold)
    return sh


def pic(slide, name, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return slide.shapes.add_picture(os.path.join(HERE, name), Inches(x), Inches(y), **kw)


def arrow(slide, x, y, w, h=0.16):
    sh = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = GREY_LINE
    sh.line.fill.background(); noshadow(sh)
    return sh


def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def drop(shape):
    shape._element.getparent().remove(shape._element)


